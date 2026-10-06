"""Generate LAPORAN-Tugas-4.docx mengikuti template docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)


def para(text="", bold=False, align=None, italic=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = "Times New Roman"
        r.font.color.rgb = RGBColor(0, 0, 0)
    return h


def code(text):
    for line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(9)


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def screenshot(text):
    p = doc.add_paragraph()
    r = p.add_run(f"[{text}]")
    r.italic = True
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


# ---------------- COVER ----------------
para("CAK3EAB3 KOMPUTASI AWAN DAN TERDISTRIBUSI", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Tugas-4", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para()
para("KOMUNIKASI PADA SISTEM TERDISTRIBUSI", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para()
para("Kelompok: Devstack", align=WD_ALIGN_PARAGRAPH.CENTER)
para("KELAS IFX-49-01", align=WD_ALIGN_PARAGRAPH.CENTER)
para("Dosen: XXX", align=WD_ALIGN_PARAGRAPH.CENTER)
para()
para("PROGRAM STUDI XXX", align=WD_ALIGN_PARAGRAPH.CENTER)
para("FAKULTAS INFORMATIKA", align=WD_ALIGN_PARAGRAPH.CENTER)
para("UNIVERSITAS TELKOM", align=WD_ALIGN_PARAGRAPH.CENTER)
para("BANDUNG", align=WD_ALIGN_PARAGRAPH.CENTER)
para("2026", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ---------------- KETENTUAN ----------------
heading("Ketentuan", 1)
for i, t in enumerate([
    "Dikerjakan secara kelompok.",
    "Dikerjakan langsung pada lembar soal.",
    "Jawaban ditulis dengan tulisan tangan yang rapi.",
    "Tulisan tangan harus terbaca dengan baik. Tulisan yang tidak terbaca diberi nilai 0.",
    "Untuk soal perhitungan, maka cara atau langkah pengerjaan wajib dituliskan.",
    "Setiap anggota kelompok wajib mengerjakan soal (ikut berkontribusi).",
    "Anggota kelompok yang tidak ikut mengerjakan maka nilainya nol, nama mahasiswa tersebut tidak perlu ditulis di tabel.",
    "Unggah jawaban ke LMS diwakili oleh salah satu anggota kelompok. Lembar fisik dikumpulkan pada saat pertemuan di kelas.",
    "Lengkapi tabel pernyataan di bawah ini sebagai syarat nilai diinput.",
], 1):
    doc.add_paragraph(f"{i}. {t}")

heading("PERNYATAAN", 1)
para("Saya menyatakan bahwa saya:")
for t in [
    "(a) benar-benar ikut mengerjakan soal",
    "(b) telah menjelaskan jawaban saya ke teman anggota kelompok",
    "(c) telah memahami penjelasan dari teman anggota kelompok",
    "(d) tidak mencontek jawaban dari kelompok lain",
    "(e) tidak memberikan jawaban PR ini ke kelompok lain",
]:
    doc.add_paragraph(t)

rows = [
    ("Nama", "NIM", "Kontribusi Tugas", "Tanda Tangan"),
    ("Muhammad Kafaby", "103012580045", "Soal 1 (RPC)", ""),
    ("Davi Pramudya Putra", "103012580056", "Soal 2 (Message Queueing)", ""),
    ("Dzaki Alwan Fitjatullah", "103012580006", "Soal 3 (Reticulum Announce)", ""),
    ("M. Faishal Rafid", "103012580034", "Deployment Caddy + dokumentasi", ""),
]
tbl = doc.add_table(rows=len(rows), cols=4)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.text = val
        if ri == 0:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True

doc.add_page_break()

# ---------------- SOAL 1 ----------------
heading("TUGAS 4", 1)
heading("1. RPC (Remote Procedure Call)", 2)
para("Soal:", bold=True)
bullets([
    "Buatlah program sederhana menggunakan RPC!",
    "Server mempunyai sebuah service/fungsi yaitu mengembalikan list bilangan prima. Parameter fungsi adalah batas terbesar angka. Contoh rpc_prima(11) maka hasilnya adalah 2, 3, 5, 7, 11.",
    "Client meminta service/fungsi yang ada pada server menggunakan RPC.",
    "Masukkan program/service pada container dan pastikan service pada container bisa diakses. Screenshot hasilnya.",
])
para("Jawaban:", bold=True)
para("RPC dibuat memakai XML-RPC (pustaka standar Python xmlrpc). Server mengekspos fungsi "
     "rpc_prima(batas) yang mengembalikan seluruh bilangan prima <= batas menggunakan algoritma "
     "Sieve of Eratosthenes.")
para("Alur kerja:")
bullets([
    "Server mendaftarkan fungsi rpc_prima pada SimpleXMLRPCServer dan listen di port 8000.",
    "Client membuka ServerProxy ke server, lalu memanggil rpc_prima(11) seperti memanggil fungsi lokal.",
    "Server menghitung bilangan prima dan mengirim balik hasilnya via HTTP.",
])
para("Kode server (soal1-rpc/prima_server.py):", bold=True)
code('''from xmlrpc.server import SimpleXMLRPCServer


def rpc_prima(batas):
    """Kembalikan list bilangan prima <= batas."""
    if batas < 2:
        return []
    sieve = [True] * (batas + 1)
    sieve[0] = sieve[1] = False
    for i in range(2, int(batas ** 0.5) + 1):
        if sieve[i]:
            for j in range(i * i, batas + 1, i):
                sieve[j] = False
    return [n for n, ok in enumerate(sieve) if ok]


if __name__ == "__main__":
    server = SimpleXMLRPCServer(("0.0.0.0", 8000), allow_none=True)
    server.register_function(rpc_prima, "rpc_prima")
    print("RPC server prima jalan di 0.0.0.0:8000", flush=True)
    server.serve_forever()''')
para("Kode client (soal1-rpc/prima_client.py):", bold=True)
code('''import sys
import xmlrpc.client

host = sys.argv[1] if len(sys.argv) > 1 else "localhost"
batas = int(sys.argv[2]) if len(sys.argv) > 2 else 11

proxy = xmlrpc.client.ServerProxy(f"http://{host}:8000")
hasil = proxy.rpc_prima(batas)
print(f"rpc_prima({batas}) = {', '.join(map(str, hasil))}", flush=True)''')
para("Langkah pengerjaan (perhitungan) untuk rpc_prima(11):", bold=True)
bullets([
    "Buat tabel boolean 0..11, tandai 0 dan 1 bukan prima.",
    "i = 2 (2^2 = 4 <= 11): tandai kelipatan 2, yaitu 4, 6, 8, 10 bukan prima.",
    "i = 3 (3^2 = 9 <= 11): tandai kelipatan 3, yaitu 6, 9 bukan prima.",
    "i = 4 (4^2 = 16 > 11): berhenti.",
    "Sisa yang bertanda prima: 2, 3, 5, 7, 11.",
])
para("Cara menjalankan (container):", bold=True)
code('''cd soal1-rpc
docker compose up --build --abort-on-container-exit''')
para("Hasil:", bold=True)
code('''rpc-prima-server  | RPC server prima jalan di 0.0.0.0:8000
rpc-prima-client  | rpc_prima(11) = 2, 3, 5, 7, 11''')
screenshot("Screenshot: output client rpc_prima(11) = 2, 3, 5, 7, 11")

doc.add_page_break()

# ---------------- SOAL 2 ----------------
heading("2. Message Queueing", 2)
para("Soal:", bold=True)
bullets([
    "Buatlah satu program sederhana menggunakan metode message queueing (MQTT / RabbitMQ / ZeroMQ).",
    "Buat program yang mengirimkan data hasil sensor (dummy) secara periodik (program mengirim data dummy setiap 1 menit sekali ke program lain).",
    "Masukkan program/service pada container dan pastikan service pada container bisa diakses. Screenshot hasilnya.",
])
para("Jawaban:", bold=True)
para("Message queueing dibuat memakai MQTT (broker eclipse-mosquitto, pustaka client paho-mqtt). "
     "Terdapat tiga komponen:")
bullets([
    "Broker (mosquitto): perantara pesan, listen port 1883.",
    "Publisher: mengirim data sensor dummy tiap 60 detik ke topik sensor/dummy.",
    "Subscriber: subscribe topik sensor/dummy dan menampilkan data yang diterima.",
])
para("Kode publisher (soal2-mqtt/publisher.py):", bold=True)
code('''import json
import os
import random
import time

import paho.mqtt.client as mqtt

BROKER = os.getenv("MQTT_BROKER", "mosquitto")
INTERVAL = int(os.getenv("INTERVAL", "60"))
TOPIC = "sensor/dummy"

client = mqtt.Client(mqtt.CallbackAPIVersion.VERSION2)
client.connect(BROKER, 1883, 60)
client.loop_start()

print(f"Publisher kirim data sensor dummy tiap {INTERVAL} detik ke topik '{TOPIC}'", flush=True)
while True:
    data = {
        "suhu": round(random.uniform(25, 35), 2),
        "kelembapan": round(random.uniform(40, 80), 2),
        "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
    }
    client.publish(TOPIC, json.dumps(data))
    print("kirim:", data, flush=True)
    time.sleep(INTERVAL)''')
para("Kode subscriber (soal2-mqtt/subscriber.py):", bold=True)
code('''import os

import paho.mqtt.client as mqtt

BROKER = os.getenv("MQTT_BROKER", "mosquitto")
TOPIC = "sensor/dummy"


def on_connect(client, userdata, flags, rc, props=None):
    print(f"Subscriber terhubung ke broker {BROKER}, subscribe '{TOPIC}'", flush=True)
    client.subscribe(TOPIC)


def on_message(client, userdata, msg):
    print(f"terima [{msg.topic}] {msg.payload.decode()}", flush=True)


client = mqtt.Client(mqtt.CallbackAPIVersion.VERSION2)
client.on_connect = on_connect
client.on_message = on_message
client.connect(BROKER, 1883, 60)
client.loop_forever()''')
para("Konfigurasi broker (soal2-mqtt/mosquitto.conf):", bold=True)
code('''listener 1883
protocol mqtt
allow_anonymous true

listener 9001
protocol websockets
allow_anonymous true''')
para("Cara menjalankan (container):", bold=True)
code('''cd soal2-mqtt
docker compose up -d --build
docker logs mqtt-publisher
docker logs mqtt-subscriber''')
para("Hasil (publisher kirim tiap 60 detik, subscriber menerima):", bold=True)
code('''mqtt-publisher   | Publisher kirim data sensor dummy tiap 60 detik ke topik 'sensor/dummy'
mqtt-publisher   | kirim: {'suhu': 32.9, 'kelembapan': 48.85, 'timestamp': '2026-10-06 07:02:12'}
mqtt-publisher   | kirim: {'suhu': 29.29, 'kelembapan': 57.42, 'timestamp': '2026-10-06 07:03:12'}
mqtt-subscriber  | Subscriber terhubung ke broker mosquitto, subscribe 'sensor/dummy'
mqtt-subscriber  | terima [sensor/dummy] {"suhu": 29.29, "kelembapan": 57.42, "timestamp": "2026-10-06 07:03:12"}''')
screenshot("Screenshot: log publisher dan subscriber")

doc.add_page_break()

# ---------------- SOAL 3 ----------------
heading("3. Reticulum Announce", 2)
para("Soal:", bold=True)
para("Buatlah program sederhana menggunakan reticulum yang bisa melakukan announce. "
     "Node A melakukan announce dan node B menerima announce tersebut kemudian menampilkannya.")
para("Jawaban:", bold=True)
para("Dibuat dua node memakai pustaka Reticulum (RNS) yang saling terhubung lewat TCP Interface "
     "(Node A sebagai TCP server, Node B sebagai TCP client). Node A membuat Destination dan memanggil "
     "announce(), Node B mendaftarkan announce handler untuk menerima dan menampilkan announce.")
para("Kode Node A (soal3-reticulum/node_a.py):", bold=True)
code('''import os
import time

import RNS

CONFIGDIR = os.getenv("RNS_CONFIG", "/root/.reticulum")

CONFIG = """[reticulum]
  enable_transport = No
  share_instance = No
  panic_on_interface_error = No

[logging]
  loglevel = 4

[interfaces]
  [[TCP Server]]
    type = TCPServerInterface
    interface_enabled = True
    listen_ip = 0.0.0.0
    listen_port = 4242
"""


def write_config():
    os.makedirs(CONFIGDIR, exist_ok=True)
    path = os.path.join(CONFIGDIR, "config")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(CONFIG)


write_config()
RNS.Reticulum(configdir=CONFIGDIR)

identity = RNS.Identity()
dest = RNS.Destination(
    identity, RNS.Destination.IN, RNS.Destination.SINGLE, "tugas", "announce"
)

print("Node A siap. Destination hash:", dest.hexhash, flush=True)
while True:
    dest.announce(app_data=b"halo dari node A")
    print("Node A mengirim announce:", dest.hexhash, flush=True)
    time.sleep(15)''')
para("Kode Node B (soal3-reticulum/node_b.py):", bold=True)
code('''import os
import time

import RNS

CONFIGDIR = os.getenv("RNS_CONFIG", "/root/.reticulum")

CONFIG = """[reticulum]
  enable_transport = No
  share_instance = No
  panic_on_interface_error = No

[logging]
  loglevel = 4

[interfaces]
  [[TCP Client]]
    type = TCPClientInterface
    interface_enabled = True
    target_host = node_a
    target_port = 4242
"""


def write_config():
    os.makedirs(CONFIGDIR, exist_ok=True)
    path = os.path.join(CONFIGDIR, "config")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(CONFIG)


write_config()
RNS.Reticulum(configdir=CONFIGDIR)


class AnnounceHandler:
    aspect_filter = "tugas.announce"

    def received_announce(self, destination_hash, announced_identity, app_data):
        print("Node B menerima announce:", RNS.hexrep(destination_hash, delimit=False), flush=True)
        if app_data:
            print("  app_data:", app_data.decode(errors="replace"), flush=True)


RNS.Transport.register_announce_handler(AnnounceHandler())
print("Node B menunggu announce dari node A...", flush=True)
while True:
    time.sleep(1)''')
para("Catatan: register_announce_handler hanya menerima handler yang punya atribut aspect_filter, "
     "sehingga handler Node B wajib mendefinisikan aspect_filter.")
para("Cara menjalankan (container):", bold=True)
code('''cd soal3-reticulum
docker compose up -d --build
docker logs rns-node-a
docker logs rns-node-b''')
para("Hasil:", bold=True)
code('''rns-node-a  | Node A siap. Destination hash: 1914016191ac77c5b5660846f8bbb646
rns-node-a  | Node A mengirim announce: 1914016191ac77c5b5660846f8bbb646
rns-node-b  | Node B menunggu announce dari node A...
rns-node-b  | Node B menerima announce: 1914016191ac77c5b5660846f8bbb646
rns-node-b  |   app_data: halo dari node A''')
screenshot("Screenshot: log Node A mengirim announce dan Node B menerima announce")

doc.add_page_break()

# ---------------- LAMPIRAN ----------------
heading("Lampiran - Deployment Production (Caddy + Domain)", 2)
para("Ketiga service di-deploy pada satu host dan diakses lewat domain menggunakan Caddy sebagai "
     "reverse proxy dengan TLS otomatis (Let's Encrypt).")
para("Domain & DNS:", bold=True)
dns_rows = [
    ("Domain", "Tipe", "Arah", "Service"),
    ("rpc.devstacklabs.net", "A", "43.157.210.87", "RPC server (rpc-prima-server:8000)"),
    ("mqtt.devstacklabs.net", "A", "43.157.210.87", "MQTT WebSocket (mqtt-broker:9001)"),
    ("kanban.devstacklabs.net", "A", "43.157.210.87", "Vikunja (existing)"),
]
t2 = doc.add_table(rows=len(dns_rows), cols=4)
t2.style = "Table Grid"
for ri, row in enumerate(dns_rows):
    for ci, val in enumerate(row):
        c = t2.cell(ri, ci)
        c.text = val
        if ri == 0:
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True
para("Network: service RPC & MQTT digabung ke network vikunja_default (external) agar dapat "
     "di-resolve oleh Caddy yang sudah berjalan.")
para("Caddyfile (/home/repo/vikunja/Caddyfile):", bold=True)
code('''kanban.devstacklabs.net {
	reverse_proxy vikunja:3456
}

rpc.devstacklabs.net {
	reverse_proxy rpc-prima-server:8000
}

mqtt.devstacklabs.net {
	reverse_proxy mqtt-broker:9001
}''')
para("MQTT diekspos melalui WebSocket (listener 9001) karena Caddy reverse_proxy mendukung WebSocket "
     "secara native, sehingga tidak perlu build plugin caddy-l4 untuk TCP mentah.")
para("Deploy:", bold=True)
code('''cd soal1-rpc && docker compose up -d --build
cd ../soal2-mqtt && docker compose up -d --build
docker exec vikunja-caddy-1 caddy reload --config /etc/caddy/Caddyfile''')
para("Verifikasi:", bold=True)
code('''python3 -c "import xmlrpc.client; print(xmlrpc.client.ServerProxy('https://rpc.devstacklabs.net').rpc_prima(11))"
# rpc_prima(11) = [2, 3, 5, 7, 11]''')
para("Hasil:", bold=True)
code('''https://rpc.devstacklabs.net    -> rpc_prima(11) = [2, 3, 5, 7, 11]
wss://mqtt.devstacklabs.net/    -> terima [sensor/dummy] {"suhu":30.98,...}
https://kanban.devstacklabs.net -> 200''')
screenshot("Screenshot: akses https://rpc.devstacklabs.net dan client MQTT via wss://mqtt.devstacklabs.net")
para("Catatan keamanan: allow_anonymous true pada broker hanya untuk demo; untuk production wajib "
     "diaktifkan autentikasi (user/password atau client certificate) beserta ACL.")

doc.save("/out/LAPORAN-Tugas-4.docx")
print("selesai: LAPORAN-Tugas-4.docx")
